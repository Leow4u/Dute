from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import unicodedata
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

PROJECT_ID = "dute-btg"
DATABASE = "(default)"
DEFAULT_CUSTOMER = "BTG Pactual"
WAREHOUSE_CODE = "DUTE-SP"
MANAGED_BY = "tools/import_btg_firestore.py"
SKU_CLIENT_ITEM_OVERRIDES = {
    # O mesmo SKU foi recadastrado no Catálogo; as linhas antigas da aba SKUs
    # pertencem ao código atual abaixo. A decisão fica explícita e auditável.
    "BTG-1033": "A.01.0802.004.P",
    "BTG-1034": "A.01.0802.004.M",
}
HTTP_TIMEOUT_SECONDS = 60
ROOT = Path(__file__).resolve().parents[1]
REAL_INVENTORY_XLSX = Path(
    r"C:\Users\leona\OneDrive - Dutelog\Pasta Profissional\Empresas\DUTELOG\Operacional\Armazenamento\BTG Pactual - Store\Armazenamento\Inventário_BTG_Consolidado.xlsx"
)
REAL_BILLING_XLSX = Path(
    r"C:\Users\leona\OneDrive - Dutelog\Pasta Profissional\Empresas\DUTELOG\Operacional\Armazenamento\BTG Pactual - Store\Armazenamento\Armazenamento BTG_Fechamento mensal.xlsx"
)
INVENTORY_XLSX = Path(
    os.environ.get(
        "BTG_INVENTORY_XLSX",
        str(
            REAL_INVENTORY_XLSX
            if REAL_INVENTORY_XLSX.exists()
            else ROOT
            / "attached_assets"
            / "Inventário_BTG_Consolidado_1783512047835.xlsx"
        ),
    )
)
BILLING_XLSX = Path(
    os.environ.get(
        "BTG_BILLING_XLSX",
        str(
            REAL_BILLING_XLSX
            if REAL_BILLING_XLSX.exists()
            else ROOT
            / "attached_assets"
            / "Armazenamento_BTG_Fechamento_mensal_1783520952576.xlsx"
        ),
    )
)
GCLOUD = (
    os.environ.get("GCLOUD")
    or shutil.which("gcloud")
    or r"C:\Users\leona\AppData\Local\Google\Cloud SDK\google-cloud-sdk\bin\gcloud.cmd"
)

INVENTORY_COLLECTIONS = (
    "products",
    "skus",
    "locations",
    "stock_balances",
    "client_inventory",
    "client_order_tracking",
    "movements",
    "orders",
)
BILLING_COLLECTIONS = (
    "billing_price_list",
    "billing_periods",
    "billing_manual",
    "billing_config",
    "client_invoices",
)


def token() -> str:
    return subprocess.check_output(
        [GCLOUD, "auth", "print-access-token"],
        text=True,
    ).strip()


def clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def number(value: Any, default: float = 0) -> float:
    value = clean(value)
    if value is None:
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def integer(value: Any, default: int = 0) -> int:
    return int(number(value, default))


def key_text(value: Any) -> str:
    value = clean(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "")


def slug(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or "sem-valor"))
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn").lower()
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")[:80]
    return text or "sem-valor"


def stable_id(prefix: str, *parts: Any) -> str:
    normalized = "\x1f".join(key_text(part) for part in parts)
    digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()[:16]
    readable = slug(parts[-1] if parts else prefix)[:48]
    return f"{prefix}_{readable}_{digest}"


def client_inventory_id(customer_code: str, sku_code: Any) -> str:
    return re.sub(
        r"[^A-Za-z0-9_.-]",
        "_",
        f"cinv_{customer_code}__{key_text(sku_code)}",
    )[:900]


def location_id_from_sku(sku: dict[str, Any]) -> str:
    rua = sku.get("rua") or "Sem rua"
    prateleira = sku.get("prateleira") or "Sem prateleira"
    caixa = sku.get("caixa") or sku.get("endereco") or "Sem caixa"
    return f"loc_{slug(rua)}_{slug(prateleira)}_{slug(caixa)}"


def row_dict(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    return {headers[i]: clean(row[i]) for i in range(min(len(headers), len(row)))}


def sheet_rows(
    ws: Any,
    header_row: int,
    *,
    include_source_row: bool = False,
) -> list[dict[str, Any]]:
    headers = [clean(cell.value) for cell in ws[header_row]]
    headers = [
        str(header) if header else f"col_{index}"
        for index, header in enumerate(headers)
    ]
    rows: list[dict[str, Any]] = []
    for source_row, row in enumerate(
        ws.iter_rows(min_row=header_row + 1, values_only=True),
        start=header_row + 1,
    ):
        item = row_dict(headers, row)
        if not any(value is not None for value in item.values()):
            continue
        if include_source_row:
            item["__source_row"] = source_row
        rows.append(item)
    return rows


def load_inventory_physical(
    inventory_xlsx: Path | str | None = None,
) -> dict[str, list[dict[str, Any]]]:
    path = Path(inventory_xlsx or INVENTORY_XLSX)
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        products: list[dict[str, Any]] = []
        for row in sheet_rows(wb["Catálogo"], 7, include_source_row=True):
            cod_cliente = key_text(row.get("CÓD. CLIENTE"))
            produto = clean(row.get("PRODUTO"))
            if not cod_cliente or not produto:
                continue
            estoque = integer(row.get("ESTOQUE BTG"))
            preco = number(row.get("PREÇO UNIT."))
            products.append(
                {
                    "source_row": row["__source_row"],
                    "customer_code": DEFAULT_CUSTOMER,
                    "client_item_code": cod_cliente,
                    "cod_cliente": cod_cliente,
                    "genero": row.get("GÊNERO"),
                    "categoria": row.get("CATEGORIA"),
                    "grupo": row.get("GRUPO"),
                    "produto": produto,
                    "cor": row.get("COR"),
                    "tamanho": row.get("TAM"),
                    "preco_unit": preco,
                    "estoque_btg": estoque,
                    "custo_total": number(row.get("CUSTO TOTAL"), estoque * preco),
                    "num_skus": integer(row.get("# SKUs")),
                }
            )

        product_price_by_code: dict[str, float] = {}
        for product in products:
            product_price_by_code.setdefault(
                key_text(product["cod_cliente"]),
                number(product.get("preco_unit")),
            )

        skus: list[dict[str, Any]] = []
        for row in sheet_rows(wb["SKUs"], 5, include_source_row=True):
            cod_btg = key_text(row.get("CÓD. BTG"))
            if not cod_btg:
                continue
            cod_cliente = SKU_CLIENT_ITEM_OVERRIDES.get(
                cod_btg,
                key_text(row.get("CÓD. CLIENTE")),
            )
            estoque_atual = integer(row.get("ESTOQUE ATUAL"))
            preco_unit = number(
                row.get("PREÇO CUSTO"),
                product_price_by_code.get(cod_cliente, 0),
            )
            skus.append(
                {
                    "source_row": row["__source_row"],
                    "customer_code": DEFAULT_CUSTOMER,
                    "sku_code": cod_btg,
                    "cod_btg": cod_btg,
                    "estoque_tipo": row.get("ESTOQUE"),
                    "client_item_code": cod_cliente or None,
                    "cod_cliente": cod_cliente or None,
                    "genero": row.get("GÊNERO"),
                    "categoria": row.get("CATEGORIA"),
                    "produto": row.get("PRODUTO"),
                    "cor": row.get("COR"),
                    "tamanho": row.get("TAM"),
                    "estoque_atual": estoque_atual,
                    "estado": row.get("ESTADO"),
                    "rua": row.get("RUA"),
                    "prateleira": row.get("PRATELEIRA"),
                    "caixa": row.get("CAIXA"),
                    "endereco": row.get("ENDEREÇO"),
                    "preco_unit": preco_unit,
                    "unit_cost": preco_unit,
                    "total_cost": estoque_atual * preco_unit,
                }
            )

        movements: list[dict[str, Any]] = []
        for row in sheet_rows(wb["Movimentações"], 5, include_source_row=True):
            cod_sku = key_text(row.get("CÓD. SKU"))
            if not cod_sku:
                continue
            movements.append(
                {
                    "source_row": row["__source_row"],
                    "customer_code": DEFAULT_CUSTOMER,
                    "data": clean(row.get("DATA")),
                    "tipo_mov": row.get("TIPO MOV."),
                    "cod_sku": cod_sku,
                    "sku_code": cod_sku,
                    "produto": row.get("PRODUTO"),
                    "qtde": integer(row.get("QTDE")),
                    "sinal": integer(row.get("SINAL (±)")),
                    "estoque_tipo": row.get("ESTOQUE"),
                    "origem": row.get("ORIGEM"),
                    "destino": row.get("DESTINO"),
                    "evento_ref": row.get("EVENTO/REF"),
                    "doc_nf": row.get("DOC/NF"),
                    "responsavel": row.get("RESPONSÁVEL"),
                    "caixa_origem": row.get("CAIXA ORIGEM"),
                }
            )

        orders: list[dict[str, Any]] = []
        for row in sheet_rows(wb["Pedidos"], 5, include_source_row=True):
            num_pedido = key_text(row.get("Nº PEDIDO"))
            if not num_pedido:
                continue
            cod_cliente = key_text(row.get("CÓD. CLIENTE"))
            orders.append(
                {
                    "source_row": row["__source_row"],
                    "customer_code": DEFAULT_CUSTOMER,
                    "order_number": num_pedido,
                    "num_pedido": num_pedido,
                    "data_solic": clean(row.get("DATA SOLIC.")),
                    "canal": row.get("CANAL"),
                    "solicitante": row.get("SOLICITANTE"),
                    "client_item_code": cod_cliente or None,
                    "cod_cliente": cod_cliente or None,
                    "product_name": row.get("PRODUTO"),
                    "produto": row.get("PRODUTO"),
                    "qtd_solic": integer(row.get("QTD SOLIC.")),
                    "qtd_sep": integer(row.get("QTD SEP.")),
                    "estoque": row.get("ESTOQUE"),
                    "status": row.get("STATUS"),
                    "data_exped": clean(row.get("DATA EXPED.")),
                    "nf": row.get("NF"),
                    "po_btg": row.get("PO BTG"),
                    "obs": row.get("OBS"),
                }
            )
    finally:
        wb.close()

    return {
        "products": products,
        "skus": skus,
        "movements": movements,
        "orders": orders,
    }


def _managed(data: dict[str, Any]) -> dict[str, Any]:
    return {
        **data,
        "managed_by": MANAGED_BY,
    }


def _stable_physical_documents(
    rows: Iterable[dict[str, Any]],
    prefix: str,
) -> list[dict[str, Any]]:
    occurrences: Counter[str] = Counter()
    documents: list[dict[str, Any]] = []
    for row in rows:
        payload = {key: value for key, value in row.items() if key != "source_row"}
        signature = json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str)
        occurrence = occurrences[signature]
        occurrences[signature] += 1
        documents.append(
            _managed(
                {
                    "id": stable_id(prefix, signature, occurrence),
                    **payload,
                    "source_row": row["source_row"],
                }
            )
        )
    return documents


def build_inventory_documents(
    physical: dict[str, list[dict[str, Any]]],
) -> dict[str, list[dict[str, Any]]]:
    product_rows_by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in physical["products"]:
        product_rows_by_code[key_text(row["cod_cliente"])].append(row)

    products: list[dict[str, Any]] = []
    for cod_cliente, rows in product_rows_by_code.items():
        first = rows[0]
        products.append(
            _managed(
                {
                    "id": stable_id("product", DEFAULT_CUSTOMER, cod_cliente),
                    "customer_code": DEFAULT_CUSTOMER,
                    "client_item_code": cod_cliente,
                    "cod_cliente": cod_cliente,
                    "genero": first.get("genero"),
                    "categoria": first.get("categoria"),
                    "grupo": first.get("grupo"),
                    "produto": first.get("produto"),
                    "cor": first.get("cor"),
                    "tamanho": first.get("tamanho"),
                    "preco_unit": number(first.get("preco_unit")),
                    # Catálogo represents one logical product. Repeated physical rows
                    # are evidence to reconcile, not quantities to add together.
                    "estoque_btg": integer(first.get("estoque_btg")),
                    "custo_total": number(first.get("custo_total")),
                    "num_skus": integer(first.get("num_skus")),
                    "physical_row_count": len(rows),
                    "source_rows": sorted(integer(row["source_row"]) for row in rows),
                    "source": "Inventário_BTG_Consolidado.xlsx :: Catálogo",
                }
            )
        )

    sku_rows_by_client: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in physical["skus"]:
        client_code = key_text(row.get("cod_cliente"))
        if client_code:
            sku_rows_by_client[client_code].append(row)
    for client_code, rows in sku_rows_by_client.items():
        if client_code in product_rows_by_code:
            continue
        first = rows[0]
        quantity = sum(integer(row.get("estoque_atual")) for row in rows)
        unit_cost = number(first.get("unit_cost"), number(first.get("preco_unit")))
        products.append(
            _managed(
                {
                    "id": stable_id("product", DEFAULT_CUSTOMER, client_code),
                    "customer_code": DEFAULT_CUSTOMER,
                    "client_item_code": client_code,
                    "cod_cliente": client_code,
                    "genero": first.get("genero"),
                    "categoria": first.get("categoria"),
                    "grupo": first.get("grupo"),
                    "produto": first.get("produto") or client_code,
                    "cor": first.get("cor"),
                    "tamanho": first.get("tamanho"),
                    "preco_unit": unit_cost,
                    "estoque_btg": quantity,
                    "custo_total": quantity * unit_cost,
                    "num_skus": len(
                        {key_text(row.get("cod_btg")) for row in rows}
                    ),
                    "physical_row_count": len(rows),
                    "source_rows": sorted(
                        integer(row["source_row"]) for row in rows
                    ),
                    "catalog_status": "synthesized_from_real_sku_rows",
                    "source": "Inventário_BTG_Consolidado.xlsx :: SKUs",
                }
            )
        )

    product_by_code = {key_text(row["cod_cliente"]): row for row in products}
    sku_rows_by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in physical["skus"]:
        sku_rows_by_code[key_text(row["cod_btg"])].append(row)

    skus: list[dict[str, Any]] = []
    for cod_btg, rows in sku_rows_by_code.items():
        first = rows[0]
        client_codes = sorted(
            {key_text(row.get("cod_cliente")) for row in rows if row.get("cod_cliente")}
        )
        client_code = client_codes[0] if len(client_codes) == 1 else None
        product = product_by_code.get(client_code or "", {})
        quantity = sum(integer(row.get("estoque_atual")) for row in rows)
        total_cost = sum(number(row.get("total_cost")) for row in rows)
        unit_cost = number(first.get("unit_cost"), number(product.get("preco_unit")))
        locations = sorted({location_id_from_sku(row) for row in rows})
        skus.append(
            _managed(
                {
                    "id": stable_id("sku", DEFAULT_CUSTOMER, cod_btg),
                    "customer_code": DEFAULT_CUSTOMER,
                    "sku_code": cod_btg,
                    "cod_btg": cod_btg,
                    "estoque_tipo": first.get("estoque_tipo"),
                    "client_item_code": client_code,
                    "client_item_codes": client_codes,
                    "cod_cliente": client_code,
                    "genero": first.get("genero"),
                    "categoria": first.get("categoria"),
                    "produto": first.get("produto"),
                    "cor": first.get("cor"),
                    "tamanho": first.get("tamanho"),
                    "estoque_atual": quantity,
                    "estado": first.get("estado"),
                    "rua": first.get("rua"),
                    "prateleira": first.get("prateleira"),
                    "caixa": first.get("caixa"),
                    "endereco": first.get("endereco"),
                    "location_ids": locations,
                    "location_count": len(locations),
                    "physical_row_count": len(rows),
                    "source_rows": sorted(integer(row["source_row"]) for row in rows),
                    "preco_unit": unit_cost,
                    "unit_cost": unit_cost,
                    "total_cost": total_cost,
                    "source": "Inventário_BTG_Consolidado.xlsx :: SKUs",
                }
            )
        )

    locations_by_id: dict[str, dict[str, Any]] = {}
    location_skus: dict[str, set[str]] = defaultdict(set)
    location_rows: dict[str, list[int]] = defaultdict(list)
    balance_rows: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for sku in physical["skus"]:
        location_id = location_id_from_sku(sku)
        quantity = integer(sku.get("estoque_atual"))
        current_location = locations_by_id.setdefault(
            location_id,
            {
                "id": location_id,
                "customer_code": DEFAULT_CUSTOMER,
                "warehouse_code": WAREHOUSE_CODE,
                "aisle": sku.get("rua") or "Sem rua",
                "shelf": sku.get("prateleira") or "Sem prateleira",
                "bin": sku.get("caixa") or sku.get("endereco") or "Sem caixa",
                "label": sku.get("endereco")
                or f"Rua {sku.get('rua') or 'Sem rua'} · Prat {sku.get('prateleira') or 'Sem prateleira'} · {sku.get('caixa') or 'Sem caixa'}",
                "qty_on_hand": 0,
                "status": "active",
                "source": "Inventário_BTG_Consolidado.xlsx :: SKUs",
            },
        )
        current_location["qty_on_hand"] += quantity
        location_skus[location_id].add(key_text(sku.get("cod_btg")))
        location_rows[location_id].append(integer(sku["source_row"]))
        balance_rows[
            (
                key_text(sku.get("cod_btg")),
                key_text(sku.get("cod_cliente")),
                location_id,
            )
        ].append(sku)

    locations: list[dict[str, Any]] = []
    for location_id, location in locations_by_id.items():
        locations.append(
            _managed(
                {
                    **location,
                    "sku_count": len(location_skus[location_id]),
                    "physical_row_count": len(location_rows[location_id]),
                    "source_rows": sorted(location_rows[location_id]),
                }
            )
        )

    stock_balances: list[dict[str, Any]] = []
    for (cod_btg, cod_cliente, location_id), rows in balance_rows.items():
        first = rows[0]
        quantity = sum(integer(row.get("estoque_atual")) for row in rows)
        reserved = sum(integer(row.get("estoque_reservado")) for row in rows)
        total_cost = sum(number(row.get("total_cost")) for row in rows)
        unit_cost = number(first.get("unit_cost"))
        stock_balances.append(
            _managed(
                {
                    "id": stable_id(
                        "balance",
                        DEFAULT_CUSTOMER,
                        cod_btg,
                        cod_cliente,
                        location_id,
                    ),
                    "customer_code": DEFAULT_CUSTOMER,
                    "warehouse_code": WAREHOUSE_CODE,
                    "sku_code": cod_btg,
                    "client_item_code": cod_cliente or None,
                    "product_name": first.get("produto"),
                    "category": first.get("categoria"),
                    "condition": first.get("estado"),
                    "location_id": location_id,
                    "location_label": locations_by_id[location_id]["label"],
                    "qty_on_hand": quantity,
                    "qty_reserved": reserved,
                    "qty_available": max(0, quantity - reserved),
                    "unit_cost": unit_cost,
                    "total_cost": total_cost,
                    "source": "Inventário_BTG_Consolidado.xlsx :: SKUs",
                    "status": "available" if quantity > 0 else "empty",
                    "physical_row_count": len(rows),
                    "source_rows": sorted(integer(row["source_row"]) for row in rows),
                }
            )
        )

    client_inventory: list[dict[str, Any]] = []
    for sku in skus:
        quantity = integer(sku.get("estoque_atual"))
        unit_cost = number(sku.get("unit_cost"), number(sku.get("preco_unit")))
        client_inventory.append(
            _managed(
                {
                    "id": client_inventory_id(
                        DEFAULT_CUSTOMER,
                        sku.get("sku_code"),
                    ),
                    "customer_code": DEFAULT_CUSTOMER,
                    "sku_code": sku.get("sku_code"),
                    "client_item_code": sku.get("client_item_code"),
                    "product_name": sku.get("produto") or sku.get("sku_code"),
                    "category": sku.get("categoria") or "Sem categoria",
                    "condition": sku.get("estado"),
                    "qty_on_hand": quantity,
                    "qty_reserved": 0,
                    "qty_available": quantity,
                    "unit_cost": unit_cost,
                    "total_cost": quantity * unit_cost,
                    "status": "active" if quantity > 0 else "empty",
                    "source": "Inventário_BTG_Consolidado.xlsx :: SKUs",
                    "published_at": datetime.now(timezone.utc).isoformat(),
                }
            )
        )

    orders_by_number: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for order in physical["orders"]:
        orders_by_number[key_text(order.get("num_pedido"))].append(order)
    client_order_tracking: list[dict[str, Any]] = []
    for order_number, rows in orders_by_number.items():
        statuses = {key_text(row.get("status")).upper() for row in rows}
        delivered = bool(statuses) and statuses <= {"ENTREGUE", "FATURADO"}
        status = "delivered" if delivered else "pending"
        status_label = "Entregue" if delivered else "Pedido recebido"
        quantity_total = sum(integer(row.get("qtd_solic")) for row in rows)
        requested_at = min(
            (key_text(row.get("data_solic")) for row in rows if row.get("data_solic")),
            default="",
        )
        delivered_at = max(
            (key_text(row.get("data_exped")) for row in rows if row.get("data_exped")),
            default="",
        )
        tracking_id = stable_id("client_order", DEFAULT_CUSTOMER, order_number)
        client_order_tracking.append(
            _managed(
                {
                    "id": tracking_id,
                    "customer_code": DEFAULT_CUSTOMER,
                    "order_number": order_number,
                    "status": status,
                    "status_label": status_label,
                    "item_count": len(rows),
                    "quantity_total": quantity_total,
                    "requested_at": requested_at or None,
                    "delivered_at": delivered_at or None,
                    "timeline": [
                        {
                            "status": status,
                            "label": status_label,
                            "occurred_at": delivered_at or requested_at or None,
                            "actor_name": "Importação da planilha real",
                        }
                    ],
                    "source": "Inventário_BTG_Consolidado.xlsx :: Pedidos",
                    "published_at": datetime.now(timezone.utc).isoformat(),
                }
            )
        )

    return {
        "products": products,
        "skus": skus,
        "locations": locations,
        "stock_balances": stock_balances,
        "client_inventory": client_inventory,
        "client_order_tracking": client_order_tracking,
        "movements": _stable_physical_documents(physical["movements"], "movement"),
        "orders": _stable_physical_documents(physical["orders"], "order"),
    }


def load_inventory(
    inventory_xlsx: Path | str | None = None,
) -> dict[str, list[dict[str, Any]]]:
    return build_inventory_documents(load_inventory_physical(inventory_xlsx))


def month_to_competencia(sheet_name: str) -> str | None:
    months = {
        "JANEIRO": "01",
        "FEVEREIRO": "02",
        "MARÇO": "03",
        "MARCO": "03",
        "ABRIL": "04",
        "MAIO": "05",
        "JUNHO": "06",
        "JULHO": "07",
        "AGOSTO": "08",
        "SETEMBRO": "09",
        "OUTUBRO": "10",
        "NOVEMBRO": "11",
        "DEZEMBRO": "12",
    }
    parts = sheet_name.split()
    if len(parts) != 2 or not parts[1].isdigit():
        return None
    month = months.get(parts[0].upper())
    if not month:
        return None
    year = int(parts[1])
    year += 2000 if year < 100 else 0
    return f"{year:04d}-{month}"


def load_billing(
    billing_xlsx: Path | str | None = None,
    *,
    as_of: date | None = None,
) -> dict[str, list[dict[str, Any]]]:
    path = Path(billing_xlsx or BILLING_XLSX)
    cutoff = (as_of or date.today()).strftime("%Y-%m")
    wb = load_workbook(path, data_only=True, read_only=True)
    price_list_by_key: dict[tuple[str, str], dict[str, Any]] = {}
    manual: list[dict[str, Any]] = []
    periods: list[dict[str, Any]] = []
    ignored_future_periods: list[str] = []

    try:
        for ws in wb.worksheets:
            competencia = month_to_competencia(ws.title)
            if not competencia:
                continue
            if competencia > cutoff:
                ignored_future_periods.append(competencia)
                continue

            rows = sheet_rows(ws, 5, include_source_row=True)
            total = 0.0
            for row in rows:
                item = clean(row.get("SERVIÇO / ITEM"))
                unidade = clean(row.get("UNIDADE"))
                if not item or not unidade:
                    continue
                qtde = number(row.get("QTDE"))
                valor = number(row.get("VALOR UNIT. LPU (R$)"))
                linha_total = number(row.get("TOTAL (R$)"), qtde * valor)
                total += linha_total
                key = (str(item), str(unidade))
                price_id = stable_id("price", DEFAULT_CUSTOMER, item, unidade)
                current_price = price_list_by_key.setdefault(
                    key,
                    {
                        "id": price_id,
                        "customer_code": DEFAULT_CUSTOMER,
                        "grupo": "armazenamento",
                        "item": item,
                        "unidade": unidade,
                        "valor_unit": valor,
                        "auto_source": (
                            "ad_valorem"
                            if "ad valorem" in str(item).lower()
                            else None
                        ),
                        "ordem": len(price_list_by_key) + 1,
                        "ativo": True,
                        "managed_by": MANAGED_BY,
                    },
                )
                # The most recent non-future sheet is the active price reference.
                current_price["valor_unit"] = valor
                manual.append(
                    _managed(
                        {
                            "competencia": competencia,
                            "customer_code": DEFAULT_CUSTOMER,
                            "price_list_id": price_id,
                            "qtde": qtde,
                            "observacao": row.get("OBSERVAÇÃO"),
                            "source_row": row["__source_row"],
                        }
                    )
                )
            periods.append(
                _managed(
                    {
                        "competencia": competencia,
                        "customer_code": DEFAULT_CUSTOMER,
                        "status": "aberto" if competencia == cutoff else "historico",
                        "payment_status": (
                            "open" if competencia == cutoff else "unverified_history"
                        ),
                        "total": total,
                        "total_planilha": total,
                        "source": f"Armazenamento BTG_Fechamento mensal.xlsx :: {ws.title}",
                    }
                )
            )
    finally:
        wb.close()

    price_list = list(price_list_by_key.values())
    if not any(item.get("auto_source") == "ad_valorem" for item in price_list):
        price_list.append(
            _managed(
                {
                    "id": stable_id(
                        "price",
                        DEFAULT_CUSTOMER,
                        "Ad Valorem sobre estoque",
                        "% estoque",
                    ),
                    "customer_code": DEFAULT_CUSTOMER,
                    "grupo": "ad_valorem",
                    "item": "Ad Valorem sobre estoque",
                    "unidade": "% estoque",
                    "valor_unit": 0,
                    "auto_source": "ad_valorem",
                    "ordem": len(price_list) + 1,
                    "ativo": True,
                }
            )
        )

    return {
        "price_list": price_list,
        "manual": manual,
        "periods": periods,
        "ignored_future_periods": ignored_future_periods,
    }


def firestore_value(value: Any) -> dict[str, Any]:
    value = clean(value)
    if value is None:
        return {"nullValue": None}
    if isinstance(value, list):
        return {"arrayValue": {"values": [firestore_value(item) for item in value]}}
    if isinstance(value, dict):
        return {
            "mapValue": {
                "fields": {
                    str(key): firestore_value(item)
                    for key, item in value.items()
                    if item is not None
                }
            }
        }
    if isinstance(value, bool):
        return {"booleanValue": value}
    if isinstance(value, int):
        return {"integerValue": str(value)}
    if isinstance(value, float):
        return {"doubleValue": value}
    return {"stringValue": str(value)}


def firestore_doc(data: dict[str, Any]) -> dict[str, Any]:
    return {"fields": {key: firestore_value(value) for key, value in data.items()}}


def doc_name(
    collection_name: str,
    doc_id: str,
    *,
    project_id: str = PROJECT_ID,
    database: str = DATABASE,
) -> str:
    return (
        f"projects/{project_id}/databases/{database}/documents/"
        f"{collection_name}/{doc_id}"
    )


def upsert(
    collection_name: str,
    doc_id: str,
    data: dict[str, Any],
    *,
    project_id: str = PROJECT_ID,
    database: str = DATABASE,
) -> dict[str, Any]:
    # updateMask preserves production-only fields that this importer does not own.
    return {
        "update": {
            "name": doc_name(
                collection_name,
                doc_id,
                project_id=project_id,
                database=database,
            ),
            **firestore_doc(data),
        },
        "updateMask": {"fieldPaths": sorted(data)},
    }


def build_inventory_writes(
    inventory: dict[str, list[dict[str, Any]]],
    *,
    project_id: str = PROJECT_ID,
    database: str = DATABASE,
) -> list[dict[str, Any]]:
    writes: list[dict[str, Any]] = []
    for collection_name in INVENTORY_COLLECTIONS:
        for row in inventory[collection_name]:
            writes.append(
                upsert(
                    collection_name,
                    key_text(row["id"]),
                    row,
                    project_id=project_id,
                    database=database,
                )
            )
    return writes


def build_billing_writes(
    billing: dict[str, list[dict[str, Any]]],
    *,
    project_id: str = PROJECT_ID,
    database: str = DATABASE,
) -> list[dict[str, Any]]:
    writes: list[dict[str, Any]] = []
    for row in billing["price_list"]:
        writes.append(
            upsert(
                "billing_price_list",
                key_text(row["id"]),
                row,
                project_id=project_id,
                database=database,
            )
        )
    for period in billing["periods"]:
        writes.append(
            upsert(
                "billing_periods",
                key_text(period["competencia"]),
                period,
                project_id=project_id,
                database=database,
            )
        )
        writes.append(
            upsert(
                "client_invoices",
                key_text(period["competencia"]),
                _managed(
                    {
                        "id": key_text(period["competencia"]),
                        "customer_code": DEFAULT_CUSTOMER,
                        "competencia": period["competencia"],
                        "status": period["status"],
                        "payment_status": period.get("payment_status"),
                        "total": period["total"],
                        "currency": "BRL",
                        "source": period.get("source"),
                        "published_at": datetime.now(timezone.utc).isoformat(),
                    }
                ),
                project_id=project_id,
                database=database,
            )
        )

    manual_by_period: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in billing["manual"]:
        manual_by_period[key_text(item["competencia"])].append(
            {key: value for key, value in item.items() if key != "competencia"}
        )
    for competencia, items in manual_by_period.items():
        writes.append(
            upsert(
                "billing_manual",
                competencia,
                _managed(
                    {
                        "competencia": competencia,
                        "customer_code": DEFAULT_CUSTOMER,
                        "items": items,
                    }
                ),
                project_id=project_id,
                database=database,
            )
        )

    writes.append(
        upsert(
            "billing_config",
            "default",
            _managed(
                {
                    "cliente_nome": DEFAULT_CUSTOMER,
                    "customer_code": DEFAULT_CUSTOMER,
                    "cliente_email": "",
                    "ad_valorem_rate": 0.005,
                }
            ),
            project_id=project_id,
            database=database,
        )
    )
    return writes


def build_write_plan(
    inventory: dict[str, list[dict[str, Any]]],
    billing: dict[str, list[dict[str, Any]]],
    *,
    project_id: str = PROJECT_ID,
    database: str = DATABASE,
) -> list[dict[str, Any]]:
    return build_inventory_writes(
        inventory,
        project_id=project_id,
        database=database,
    ) + build_billing_writes(
        billing,
        project_id=project_id,
        database=database,
    )


def fetch_collection_raw(
    collection_name: str,
    *,
    project_id: str = PROJECT_ID,
    database: str = DATABASE,
    access_token: str | None = None,
) -> list[dict[str, Any]]:
    access_token = access_token or token()
    base_url = (
        f"https://firestore.googleapis.com/v1/projects/{project_id}/"
        f"databases/{database}/documents/{urllib.parse.quote(collection_name, safe='/')}"
    )
    documents: list[dict[str, Any]] = []
    page_token: str | None = None
    while True:
        query = {"pageSize": "300"}
        if page_token:
            query["pageToken"] = page_token
        request = urllib.request.Request(
            f"{base_url}?{urllib.parse.urlencode(query)}",
            headers={
                "Authorization": f"Bearer {access_token}",
                "x-goog-user-project": project_id,
            },
            method="GET",
        )
        with urllib.request.urlopen(request, timeout=HTTP_TIMEOUT_SECONDS) as response:
            payload = json.load(response)
        documents.extend(payload.get("documents", []))
        page_token = payload.get("nextPageToken")
        if not page_token:
            return documents


def create_backup(
    backup_path: Path | str,
    collections: Iterable[str],
    *,
    project_id: str = PROJECT_ID,
    database: str = DATABASE,
    access_token: str | None = None,
) -> Path:
    path = Path(backup_path).expanduser().resolve()
    if path.exists():
        raise FileExistsError(f"Backup já existe; recusando sobrescrever: {path}")
    access_token = access_token or token()
    collection_data = {
        collection: fetch_collection_raw(
            collection,
            project_id=project_id,
            database=database,
            access_token=access_token,
        )
        for collection in collections
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "format": "firestore-rest-backup-v1",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "project_id": project_id,
        "database": database,
        "collections": collection_data,
    }
    with path.open("x", encoding="utf-8", newline="\n") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
        handle.write("\n")
        handle.flush()
        os.fsync(handle.fileno())
    return path


def writes_with_backup_preconditions(
    writes: list[dict[str, Any]],
    backup_path: Path | str,
) -> list[dict[str, Any]]:
    with Path(backup_path).open(encoding="utf-8") as handle:
        backup = json.load(handle)
    existing_by_name = {
        document["name"]: document
        for documents in backup.get("collections", {}).values()
        for document in documents
        if isinstance(document, dict) and document.get("name")
    }
    guarded: list[dict[str, Any]] = []
    for write in writes:
        document_name = write.get("update", {}).get("name") or write.get("delete")
        existing = existing_by_name.get(document_name)
        if existing and existing.get("updateTime"):
            precondition = {"updateTime": existing["updateTime"]}
        elif existing:
            precondition = {"exists": True}
        else:
            precondition = {"exists": False}
        guarded.append({**write, "currentDocument": precondition})
    return guarded


def build_stale_inventory_deletes(
    inventory: dict[str, list[dict[str, Any]]],
    backup_path: Path | str,
) -> list[dict[str, Any]]:
    with Path(backup_path).open(encoding="utf-8") as handle:
        backup = json.load(handle)

    exact_collections = {
        "products",
        "skus",
        "locations",
        "stock_balances",
        "client_inventory",
        "client_order_tracking",
    }
    desired = {
        collection: {key_text(row["id"]) for row in inventory[collection]}
        for collection in exact_collections
    }
    deletes: list[dict[str, Any]] = []
    for collection in (*exact_collections, "movements", "orders"):
        for document in backup.get("collections", {}).get(collection, []):
            name = key_text(document.get("name"))
            if not name:
                continue
            document_id = name.rsplit("/", 1)[-1]
            if collection in exact_collections:
                stale = document_id not in desired[collection]
            else:
                # Importações legadas usavam IDs numéricos. Movimentos/pedidos
                # operacionais novos (mov_*/ord_*) são preservados.
                stale = document_id.isdigit()
            if not stale:
                continue
            write: dict[str, Any] = {"delete": name}
            if document.get("updateTime"):
                write["currentDocument"] = {
                    "updateTime": document["updateTime"]
                }
            else:
                write["currentDocument"] = {"exists": True}
            deletes.append(write)
    return deletes


def batch_write(
    writes: list[dict[str, Any]],
    *,
    project_id: str = PROJECT_ID,
    database: str = DATABASE,
    access_token: str | None = None,
) -> None:
    if not writes:
        return
    access_token = access_token or token()
    url = (
        f"https://firestore.googleapis.com/v1/projects/{project_id}/databases/"
        f"{database}/documents:batchWrite"
    )
    headers = {
        "Authorization": f"Bearer {access_token}",
        "x-goog-user-project": project_id,
        "Content-Type": "application/json",
    }
    for index in range(0, len(writes), 400):
        chunk = writes[index : index + 400]
        payload = json.dumps({"writes": chunk}).encode()
        request = urllib.request.Request(url, data=payload, headers=headers, method="POST")
        with urllib.request.urlopen(request, timeout=HTTP_TIMEOUT_SECONDS) as response:
            result = json.load(response)
        statuses = result.get("status", [])
        failures = [status for status in statuses if status.get("code", 0) != 0]
        if failures:
            raise RuntimeError(
                "Firestore recusou parte do batchWrite: "
                + json.dumps(failures, ensure_ascii=False)
            )


def confirmation_phrase(
    project_id: str = PROJECT_ID,
    database: str = DATABASE,
) -> str:
    return f"APPLY {project_id}/{database}"


def source_mapping_summary(
    physical: dict[str, list[dict[str, Any]]],
) -> dict[str, int]:
    product_counts = Counter(key_text(row.get("cod_cliente")) for row in physical["products"])
    catalog_codes = {code for code in product_counts if code}
    sku_to_clients: dict[str, set[str]] = defaultdict(set)
    for row in physical["skus"]:
        if row.get("cod_cliente"):
            sku_to_clients[key_text(row.get("cod_btg"))].add(
                key_text(row.get("cod_cliente"))
            )
    sku_client_codes = {
        client_code
        for client_codes in sku_to_clients.values()
        for client_code in client_codes
    }
    movement_skus = {
        key_text(row.get("cod_sku")) for row in physical["movements"]
    }
    sku_codes = set(sku_to_clients)
    return {
        "duplicate_catalog_codes": sum(count > 1 for count in product_counts.values()),
        "ambiguous_sku_mappings": sum(
            len(client_codes) > 1 for client_codes in sku_to_clients.values()
        ),
        "sku_product_codes_synthesized_from_real_skus": len(
            sku_client_codes - catalog_codes
        ),
        "historical_movement_skus_missing_from_current_skus": len(
            movement_skus - sku_codes
        ),
    }


def parse_iso_date(value: str) -> date:
    try:
        return date.fromisoformat(value)
    except ValueError as exc:
        raise argparse.ArgumentTypeError("use YYYY-MM-DD") from exc


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Prepara importação BTG; dry-run é o modo padrão.",
    )
    parser.add_argument("--inventory-xlsx", type=Path, default=INVENTORY_XLSX)
    parser.add_argument("--billing-xlsx", type=Path, default=BILLING_XLSX)
    parser.add_argument("--as-of", type=parse_iso_date, default=date.today())
    parser.add_argument("--project-id", default=PROJECT_ID)
    parser.add_argument("--database", default=DATABASE)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--dry-run", action="store_true", help="não acessa o Firestore")
    mode.add_argument("--apply", action="store_true", help="aplica o plano após backup")
    parser.add_argument("--backup", type=Path, help="arquivo JSON novo e obrigatório no apply")
    parser.add_argument("--confirm", help="frase explícita exibida no dry-run")
    return parser


def main(argv: list[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    expected_confirmation = confirmation_phrase(args.project_id, args.database)
    if args.apply:
        parser.error(
            "--apply foi congelado para o importador legado; "
            "use tools/reconcile_btg_operational_data.py com controle de reconciliação ativo"
        )
        if args.backup is None:
            parser.error("--apply exige --backup apontando para um arquivo novo")
        if args.confirm != expected_confirmation:
            parser.error(f"--apply exige --confirm {expected_confirmation!r}")
        if args.backup.expanduser().exists():
            parser.error("--backup deve apontar para um arquivo que ainda não existe")

    missing = [
        str(path)
        for path in (args.inventory_xlsx, args.billing_xlsx)
        if not path.exists()
    ]
    if missing:
        print(
            json.dumps(
                {"error": "Planilhas de origem não encontradas.", "missing": missing},
                ensure_ascii=False,
            ),
            file=sys.stderr,
        )
        return 1

    physical = load_inventory_physical(args.inventory_xlsx)
    inventory = build_inventory_documents(physical)
    billing = load_billing(args.billing_xlsx, as_of=args.as_of)
    writes = build_write_plan(
        inventory,
        billing,
        project_id=args.project_id,
        database=args.database,
    )
    summary: dict[str, Any] = {
        "mode": "apply" if args.apply else "dry-run",
        "project_id": args.project_id,
        "database": args.database,
        "as_of": args.as_of.isoformat(),
        "customer_code": DEFAULT_CUSTOMER,
        "physical_rows": {
            collection: len(rows) for collection, rows in physical.items()
        },
        "logical_documents": {
            collection: len(rows) for collection, rows in inventory.items()
        },
        "billing": {
            "price_list": len(billing["price_list"]),
            "periods": len(billing["periods"]),
            "manual_items": len(billing["manual"]),
            "ignored_future_periods": billing["ignored_future_periods"],
        },
        "mapping": source_mapping_summary(physical),
        "planned_writes": len(writes),
        "confirmation_required": expected_confirmation,
        "backup": None,
        "writes_applied": 0,
        "apply_status": "not-requested",
    }

    if args.apply:
        critical_mapping_count = (
            summary["mapping"]["ambiguous_sku_mappings"]
        )
        if critical_mapping_count:
            summary["apply_status"] = "blocked-critical-source-mappings"
            print(json.dumps(summary, ensure_ascii=False, indent=2))
            return 2
        access_token = token()
        backup = create_backup(
            args.backup,
            (*INVENTORY_COLLECTIONS, *BILLING_COLLECTIONS),
            project_id=args.project_id,
            database=args.database,
            access_token=access_token,
        )
        stale_deletes = build_stale_inventory_deletes(inventory, backup)
        guarded_writes = [
            *stale_deletes,
            *writes_with_backup_preconditions(writes, backup),
        ]
        batch_write(
            guarded_writes,
            project_id=args.project_id,
            database=args.database,
            access_token=access_token,
        )
        summary["backup"] = str(backup)
        summary["writes_applied"] = len(guarded_writes)
        summary["stale_documents_deleted"] = len(stale_deletes)
        summary["apply_status"] = "applied"

    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
